import os
import PyPDF2
import docx
import io
import uuid
import re
import json
import threading
import logging
import time
from datetime import datetime, timedelta
from flask import (Blueprint, request, render_template, jsonify, send_file,
                   redirect, url_for, flash, current_app, session, abort)
from werkzeug.utils import secure_filename
import zipfile
from werkzeug.exceptions import RequestEntityTooLarge
from sqlalchemy import desc, or_, func
import pytz
import difflib

# --- NEW IMPORTS FOR INTERNET SEARCH ---
import requests
from bs4 import BeautifulSoup
from googlesearch import search as google_search
# --- END NEW IMPORTS ---

from .extensions import db
from .models import Task, User, Wallet, Transaction, Conversation, ChatMessage, ReferralLedger, ReferralWithdrawal
from .payment_service import (initiate_payhero_payment, WORDS_PER_UNIT,
                              PRICE_PER_UNIT, ESSAY_PRICE_PER_UNIT)
from .document_formatter import create_formatted_document
from .humanizer_service.utils import count_words
from .humanizer_service.exceptions import AllModelsFailedError
from .auth_routes import auth_required
from .humanizer_service.api_clients import call_single_openai_model
from .humanizer_service.api_clients import openai_client
from .tasks import process_playground_chat_task, process_humanization_task

main_bp = Blueprint('main', __name__)

@main_bp.before_request
def redirect_heroku_to_custom_domain():
    """Redirect from the Heroku domain to the custom domain."""
    if 'herokuapp' in request.host:
        new_url = f"https://www.humanizeyouressay.online{request.full_path}"
        return redirect(new_url, code=301)

@main_bp.before_request
def redirect_non_www_to_www():
    """Redirect non-www requests to www for canonicalization."""
    if not request.host.startswith('www.') and 'herokuapp' not in request.host:
        new_host = f"www.{request.host}"
        new_url = f"{request.scheme}://{new_host}{request.full_path}"
        return redirect(new_url, code=301)


KENYAN_PHONE_REGEX = re.compile(r"^(0|\+254)7[0-9]{8}$")
logger = logging.getLogger(__name__)

WHATSAPP_NUMBER = "254712345678"
MAX_WORDS = 1925
NAIROBI_TZ = pytz.timezone('Africa/Nairobi')

# --- Security: Whitelist for allowed initial uploads ---
ALLOWED_EXTENSIONS = {'pdf', 'docx', 'zip', 'png', 'jpg', 'jpeg', 'webp', 'gif', 'txt'}

# --- Security: Blocklist for executable/dangerous files ---
BLOCKED_EXTENSIONS = {
    'exe', 'dll', 'so', 'bat', 'sh', 'ps1', 'vbs', 'msi', 'jar', 'cmd', 'scr'
}

def is_extension_allowed(filename):
    if '.' not in filename:
        return False
    ext = filename.rsplit('.', 1)[1].lower()
    if ext in BLOCKED_EXTENSIONS:
        return False
    if ext in ALLOWED_EXTENSIONS:
        return True
    return False

# --- Security: Rate Limiting ---
def check_rate_limit(limit=10, per=timedelta(minutes=1)):
    now = time.time()
    uploads = session.get('upload_timestamps', [])
    recent_uploads = [ts for ts in uploads if now - ts < per.total_seconds()]
    if len(recent_uploads) >= limit:
        return False
    recent_uploads.append(now)
    session['upload_timestamps'] = recent_uploads
    return True

def sanitize_filename(filename):
    if not filename:
        return "document"
    filename = filename.strip().replace(' ', '_')
    return re.sub(r'(?u)[^-\w.]', '', filename)

def get_display_name(user):
    if user and user.name:
        return user.name.split(' ')[0]
    elif user and user.email:
        return user.email.split('@')[0]
    return "User"

@main_bp.route("/")
def index():
    if 'user' in session:
        user = User.query.get(session['user']['uid'])
        if not user:
            session.clear()
            return redirect(url_for('main.index'))
        
        if user.is_admin:
            return redirect(url_for('admin_dashboard.dashboard'))
        
        return redirect(url_for('main.playground'))
    else:
        firebase_config = {
            "apiKey": os.getenv("FIREBASE_API_KEY"), "authDomain": os.getenv("FIREBASE_AUTH_DOMAIN"),
            "projectId": os.getenv("FIREBASE_PROJECT_ID"), "storageBucket": os.getenv("FIREBASE_STORAGE_BUCKET"),
            "messagingSenderId": os.getenv("FIREBASE_MESSAGING_SENDER_ID"), "appId": os.getenv("FIREBASE_APP_ID"),
            "measurementId": os.getenv("FIREBASE_MEASUREMENT_ID")
        }
        return render_template("home.html",
                               username=None,
                               WORDS_PER_UNIT=WORDS_PER_UNIT,
                               PRICE_PER_UNIT=PRICE_PER_UNIT + 10,
                               ESSAY_PRICE_PER_UNIT=ESSAY_PRICE_PER_UNIT + 10,
                               firebase_config=firebase_config)

@main_bp.app_errorhandler(413)
@main_bp.app_errorhandler(RequestEntityTooLarge)
def handle_file_too_large(e):
    if request.path == url_for('main.extract_text_from_file'):
        return jsonify(error="File is too large. Please upload files under 10MB."), 413
    return "File is too large.", 413

@main_bp.route('/api/get-models')
def get_models():
    models_config = current_app.config.get('MODELS', {})
    model_list = [
        {
            'value': key,
            'label': details['label'],
            'price': details['price_per_word'],
            'selected': key == current_app.config.get('DEFAULT_MODEL', 'gpt-5')
        }
        for key, details in models_config.items()
    ]
    return jsonify(model_list)

@main_bp.route("/transaction-history", methods=['GET'])
@auth_required
def transaction_history():
    user = User.query.get_or_404(session['user']['uid'])
    if not user.wallet:
        flash("No wallet found for this user.", "error")
        return redirect(url_for('main.index'))

    page = request.args.get('page', 1, type=int)
    per_page = 20

    pagination = Transaction.query.filter_by(wallet_id=user.wallet.id).order_by(desc(Transaction.created_at)).paginate(page=page, per_page=per_page, error_out=False)

    transactions = pagination.items

    if transactions:
        current_balance = user.wallet.balance

        sum_of_newer_transactions = db.session.query(db.func.sum(Transaction.amount)).filter(
            Transaction.wallet_id == user.wallet.id,
            Transaction.created_at > transactions[0].created_at
        ).scalar() or 0.0

        balance_after = current_balance - sum_of_newer_transactions

        for t in transactions:
            t.balance_after = balance_after
            balance_after -= t.amount

    for t in transactions:
        if hasattr(t, 'created_at') and t.created_at:
                t.created_at_local_str = t.created_at.replace(tzinfo=pytz.utc).astimezone(pytz.timezone('Africa/Nairobi')).strftime('%Y-%m-%d %H:%M')
        else:
                t.created_at_local_str = 'N/A'

        t.task = None
        t.display_title = t.description or 'N/A'

        if t.transaction_type == 'humanize_fee' and t.description:
            task = Task.query.filter_by(external_reference=t.description).first()
            if task and task.status == 'humanization_complete':
                t.task = task
                t.display_title = f"Humanizer Task: {task.external_reference}"


        elif t.transaction_type == 'playground_fee' and t.description and t.description.startswith('CONV-'):
            try:
                conv_id_part = t.description.split(':')[0]
                conv_id_str = conv_id_part.split('-')[1]
                conv_id = int(conv_id_str)

                conversation = Conversation.query.get(conv_id)

                if conversation and conversation.title:
                    t.display_title = f"{t.description} - {conversation.title}"
                else:
                    t.display_title = t.description
            except (IndexError, ValueError):
                pass

    username = get_display_name(user)
    return render_template("transaction_history.html",
                           user=user,
                           transactions=transactions,
                           username=username,
                           pagination=pagination,
                           wallet_balance=user.wallet.balance,
                           is_admin=user.is_admin)

@main_bp.route("/referrals")
@auth_required
def referrals_page():
    """
    Displays the user's referral statistics, leaderboard, and direct referrals.
    """
    user = User.query.get_or_404(session['user']['uid'])
    username = get_display_name(user)

    # Lifetime earnings (commissions + signup bonuses)
    lifetime_earnings = (
        db.session.query(func.sum(Transaction.amount))
        .filter(
            Transaction.wallet_id == user.wallet.id,
            Transaction.transaction_type.in_(["referral_commission", "referral_signup_bonus"])
        )
        .scalar()
        or 0.0
    )

    # Withdrawable balance = referral commissions - withdrawals
    referral_earnings = (
        db.session.query(func.sum(Transaction.amount))
        .filter(
            Transaction.wallet_id == user.wallet.id,
            Transaction.transaction_type == "referral_commission"
        )
        .scalar()
        or 0.0
    )

    withdrawn_amount = (
        db.session.query(func.sum(Transaction.amount))
        .filter(
            Transaction.wallet_id == user.wallet.id,
            Transaction.transaction_type == "withdrawal"
        )
        .scalar()
        or 0.0
    )

    withdrawable_balance = referral_earnings + withdrawn_amount  # withdrawn is negative

    # Referral counts by level (5 levels)
    referral_counts = {}
    for i in range(5):
        level = i + 1
        count = User.query.filter(User.upline.like(f"%{user.id}%")).all()
        level_count = 0
        for referred_user in count:
            if referred_user.upline:
                upline_ids = referred_user.upline.split(',')
                if len(upline_ids) > i and upline_ids[i] == str(user.id):
                    level_count += 1
        referral_counts[level] = level_count

    # Direct referrals with pagination
    page = request.args.get('page', 1, type=int)
    per_page = 10
    referrals = User.query.filter_by(referred_by=user.id).paginate(
        page=page, per_page=per_page, error_out=False
    )

    # Leaderboard filter
    filter_option = request.args.get('filter', 'all')  # 'month' or 'all'
    leaderboard_page = request.args.get('lb_page', 1, type=int)
    leaderboard_per_page = 20

    leaderboard_query = (
        db.session.query(
            User.id,
            User.name.label("display_name"),
            func.sum(Transaction.amount).label("total_earnings")
        )
        .join(Wallet, Wallet.user_id == User.id)
        .join(Transaction, Transaction.wallet_id == Wallet.id)
        .filter(Transaction.transaction_type.in_(["referral_commission", "referral_signup_bonus"]))
    )

    if filter_option == 'month':
        start_of_month = datetime(datetime.now().year, datetime.now().month, 1)
        leaderboard_query = leaderboard_query.filter(Transaction.created_at >= start_of_month)

    leaderboard_query = (
        leaderboard_query
        .group_by(User.id, User.name)
        .order_by(func.sum(Transaction.amount).desc())
    )

    leaderboard_pagination = leaderboard_query.paginate(
        page=leaderboard_page, per_page=leaderboard_per_page, error_out=False
    )

    leaderboard = [
        ((leaderboard_per_page * (leaderboard_page - 1)) + idx + 1, row.display_name, row.total_earnings)
        for idx, row in enumerate(leaderboard_pagination.items)
    ]

    # Determine user's rank
    user_rank = None
    all_leaderboard = leaderboard_query.all()
    for idx, row in enumerate(all_leaderboard, start=1):
        if row.id == user.id:
            user_rank = idx
            break

    return render_template(
        "referrals.html",
        user=user,
        username=username,
        wallet_balance=user.wallet.balance,
        lifetime_earnings=lifetime_earnings,
        withdrawable_balance=withdrawable_balance,
        is_admin=user.is_admin,
        referral_counts=referral_counts,
        referrals=referrals,
        leaderboard=leaderboard,
        user_rank=user_rank,
        page=leaderboard_page,
        total_pages=leaderboard_pagination.pages,
        filter_option=filter_option
    )


@main_bp.route("/playground_chat", methods=['POST'])
@auth_required
def playground_chat():
    user = User.query.get_or_404(session['user']['uid'])
    data = request.get_json() or {}
    input_parts = data.get('input', [])
    model_key = data.get('model')
    conversation_id = data.get('conversation_id')

    if not isinstance(input_parts, list) or not input_parts:
        return jsonify({"error": "No input provided."}), 400

    user_prompt_text = ""
    text_part_index = -1
    for i, part in enumerate(input_parts):
        if part.get("type") == "input_text":
            user_prompt_text = part.get("text", "").strip()
            text_part_index = i
            break

    # +++ START OF RE-ADDED CODE +++
    # Word count validation
    word_count = len(user_prompt_text.split())
    if word_count > 1000:
        return jsonify({'error': 'Your message exceeds the 1000-word limit.'}), 400
    # +++ END OF RE-ADDED CODE +++

    if user_prompt_text.lower().startswith('/search '):
        search_query = user_prompt_text[8:]
        try:
            search_results = list(google_search(search_query, num_results=3, sleep_interval=2))
            
            if search_results:
                top_url = search_results[0]
                verbatim_snippet = ""
                
                try:
                    headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'}
                    response = requests.get(top_url, timeout=10, headers=headers)
                    response.raise_for_status()
                    
                    soup = BeautifulSoup(response.content, 'html.parser')
                    paragraphs = soup.find_all('p')
                    snippet_parts = [p.get_text() for p in paragraphs[:5]]
                    verbatim_snippet = "\n".join(snippet_parts).strip()

                except requests.RequestException as e:
                    logger.error(f"Failed to fetch URL {top_url}: {e}")
                    verbatim_snippet = "Could not retrieve content from the source."

                if text_part_index != -1:
                    original_prompt = input_parts[text_part_index]['text']
                    
                    final_prompt = (
                        f"Please answer the user's query based on the provided search result. "
                        f"First, provide your own assessment, and then include the verbatim text snippet from the source. "
                        f"**Crucially, you must include the source URL in your assessment.**\n\n"
                        f"--- Source URL ---\n{top_url}\n\n"
                        f"--- Verbatim Snippet from Source ---\n\"{verbatim_snippet}\"\n\n"
                        f"--- User's Original Prompt ---\n{original_prompt}"
                    )
                    input_parts[text_part_index]['text'] = final_prompt

        except Exception as e:
            logger.error(f"Internet search failed for query '{search_query}': {e}", exc_info=True)
            pass

    try:
        if conversation_id:
            conversation = Conversation.query.filter_by(id=conversation_id, user_id=user.id).first_or_404()
        else:
            raw_title = " ".join(
                p.get("text", "").strip()
                for p in input_parts
                if p.get("type") == "input_text"
            ).strip()
            title = (raw_title[:60] or "New chat")
            conversation = Conversation(user_id=user.id, title=title, model_key=model_key)
            db.session.add(conversation)
            db.session.flush()
            conversation_id = conversation.id

        db.session.add(ChatMessage(
            conversation_id=conversation_id,
            role="user",
            content=json.dumps(input_parts)
        ))

        task_ref = f"PLAYGROUND-{uuid.uuid4()}"
        task_data = {
            "user_input_parts": input_parts,
            "model_key": model_key,
            "conversation_id": conversation_id,
        }

        new_task = Task(
            user_id=user.id,
            external_reference=task_ref,
            status='queued',
            diff_data=task_data,
            input_text="Playground Chat",
        )
        db.session.add(new_task)
        db.session.commit()

        process_playground_chat_task.delay(new_task.id)

        return jsonify({
            "status": "processing_started",
            "task_id": new_task.id,
            "conversation_id": conversation_id
        })

    except Exception as e:
        db.session.rollback()
        logger.error(f"Playground chat submission error for user {user.id}: {e}", exc_info=True)
        return jsonify({"error": "Failed to start the chat task."}), 500

@main_bp.route("/get_playground_task_status/<int:task_id>", methods=['GET'])
@auth_required
def get_playground_task_status(task_id):
    task = Task.query.get_or_404(task_id)

    if task.user_id != session['user']['uid']:
        return jsonify({"error": "Unauthorized"}), 403

    if task.status == 'humanization_complete':
        final_conversation_id = (task.diff_data or {}).get("final_conversation_id")
        conversation = Conversation.query.get(final_conversation_id) if final_conversation_id else None
        return jsonify({
            "status": "complete",
            "response": task.humanized_output,
            "conversation_id": final_conversation_id,
            "conversation_title": conversation.title if conversation else "Chat"
        })
    elif task.status == 'humanization_failed':
        return jsonify({
            "status": "failed",
            "error": task.error_message
        })
    else:
        return jsonify({"status": "processing"})

@main_bp.route("/humanizer", methods=['GET'])
@auth_required
def humanizer_page():
    user = User.query.get_or_404(session['user']['uid'])
    if user.is_admin:
        return redirect(url_for('admin_dashboard.dashboard'))
    if not user.wallet:
        db.session.add(Wallet(user=user, balance=0.0))
        db.session.commit()
    username = get_display_name(user)
    firebase_config = {"apiKey": os.getenv("FIREBASE_API_KEY"), "authDomain": os.getenv("FIREBASE_AUTH_DOMAIN"), "projectId": os.getenv("FIREBASE_PROJECT_ID"), "storageBucket": os.getenv("FIREBASE_STORAGE_BUCKET"), "messagingSenderId": os.getenv("FIREBASE_MESSAGING_SENDER_ID"), "appId": os.getenv("FIREBASE_APP_ID"), "measurementId": os.getenv("FIREBASE_MEASUREMENT_ID")}
    return render_template('dashboard.html', user=user, username=username, wallet_balance=user.wallet.balance, is_admin=user.is_admin, WORDS_PER_UNIT=WORDS_PER_UNIT, PRICE_PER_UNIT=PRICE_PER_UNIT, ESSAY_PRICE_PER_UNIT=ESSAY_PRICE_PER_UNIT, firebase_config=firebase_config)

@main_bp.route("/playground", methods=['GET'])
@auth_required
def playground():
    user = User.query.get_or_404(session['user']['uid'])
    username = get_display_name(user)
    wallet_balance = user.wallet.balance if user.wallet else 0
    models = current_app.config['MODELS']
    return render_template(
        "playground.html",
        user=user,
        username=username,
        wallet_balance=wallet_balance,
        is_admin=user.is_admin,
        models=models
    )

def _initiate_payment_in_background(app, amount, phone_number, transaction_ref):
    with app.app_context():
        try:
            initiate_payhero_payment(amount, phone_number, transaction_ref)
        except Exception as e:
            logger.error(f"Background payment initiation failed for {transaction_ref}: {e}")
            task = Task.query.filter_by(external_reference=transaction_ref).first()
            if task:
                task.status = 'topup_failed'
                task.error_message = "Failed to communicate with payment gateway."
                db.session.commit()

@main_bp.route("/wallet/topup", methods=["POST"])
@auth_required
def wallet_topup():
    data = request.get_json()
    amount, phone_number = data.get("amount"), data.get("phone_number")
    if not all([amount, phone_number, isinstance(amount, (int, float)), amount > 0]):
        return jsonify({"status": "error", "message": "Valid amount and phone number are required."}), 400
    if not KENYAN_PHONE_REGEX.match(phone_number):
        return jsonify({"status": "error", "message": "Invalid Kenyan phone number format."}), 400
    transaction_ref = f"TOPUP-{uuid.uuid4()}"
    temp_task = Task(external_reference=transaction_ref, input_text=f"Top-up of {amount}", phone_number=phone_number,
                     amount_charged=float(amount), status='topup_pending', user_id=session['user']['uid'])
    db.session.add(temp_task)
    db.session.commit()
    app = current_app._get_current_object()
    threading.Thread(target=_initiate_payment_in_background, args=(app, float(amount), phone_number, transaction_ref)).start()
    return jsonify({"status": "payment_initiated", "transaction_id": transaction_ref})

@main_bp.route("/wallet/status/<transaction_id>", methods=["GET"])
@auth_required
def get_topup_status(transaction_id):
    task = Task.query.filter_by(external_reference=transaction_id, user_id=session['user']['uid']).first_or_404()
    if task.status == 'topup_complete':
        return jsonify({"status": "complete", "message": "Top-up successful!"})
    elif task.status == 'topup_failed':
        return jsonify({"status": "failed", "message": task.error_message or "Top-up was cancelled or failed."})
    else:
        return jsonify({"status": "pending", "message": "Awaiting payment confirmation..."})

@main_bp.route("/submit-humanize-task", methods=["POST"])
def submit_humanize_task():
    if 'user' in session:
        return _handle_logged_in_humanization()
    else:
        return _handle_anonymous_humanization()

def _handle_logged_in_humanization():
    user = User.query.get_or_404(session['user']['uid'])
    try:
        data = request.form
        input_text = data.get("input_text", "").strip()
        word_count = count_words(input_text)
        humanizer_mode = data.get("humanizer_mode", "standard")

        if not input_text:
            return jsonify({"status": "error", "message": "Text input is required."}), 400

        if humanizer_mode != 'document_formatter' and word_count > MAX_WORDS:
            return jsonify({"status": "error", "message": f"Input exceeds the maximum of {MAX_WORDS} words."}), 413

        price = ESSAY_PRICE_PER_UNIT if humanizer_mode == 'essay' else PRICE_PER_UNIT
        amount_to_charge = round((word_count / WORDS_PER_UNIT) * price, 2)
        if user.wallet.balance < amount_to_charge:
            return jsonify({"status": "error", "message": "Your wallet balance is insufficient. Please top up your account."}), 402

        transaction_ref = f"TASK-{uuid.uuid4()}"
        user.wallet.balance -= amount_to_charge
        db.session.add(Transaction(wallet_id=user.wallet.id, amount=-amount_to_charge,
                                   transaction_type='humanize_fee', description=transaction_ref))

        # --- REFERRAL COMMISSION LOGIC START ---
        if user.upline:
            commission_levels = {
                0: 0.10,  # Level 1: 10%
                1: 0.05,  # Level 2: 5%
                2: 0.02,  # Level 3: 2%
                3: 0.01,  # Level 4: 1%
                4: 0.005  # Level 5: 0.5%
            }
            
            upline_ids = user.upline.split(',')
            
            for i, referrer_id in enumerate(upline_ids):
                if i < len(commission_levels):
                    referrer = User.query.get(referrer_id)
                    if referrer and referrer.wallet:
                        commission_amount = amount_to_charge * commission_levels[i]
                        referrer.wallet.balance += commission_amount
                        
                        # after creating a referral commission transaction...
                        commission_txn = Transaction(
                            wallet=referrer.wallet,
                            amount=commission_amount,
                            transaction_type="referral_commission",
                            description=f"Level {i+1} referral commission from {user.email}"
                        )
                        db.session.add(commission_txn)
                        db.session.flush()  # ensures commission_txn.id is available
                        
                        ledger_entry = ReferralLedger(
                            user=referrer,
                            amount=commission_amount,
                            transaction_id=commission_txn.id
                        )
                        db.session.add(ledger_entry)
        # --- REFERRAL COMMISSION LOGIC END ---

        output_format = data.get("output_format", "plain_text")
        bibtex_input = data.get("bibtex_input", "")
        english_dialect = data.get("english_dialect", "american")

        new_task = Task(
            external_reference=transaction_ref,
            input_text=input_text,
            phone_number="N/A",
            amount_charged=amount_to_charge,
            status='queued',
            user_id=user.id,
            diff_data={key: data.get(key, "") for key in ["output_format", "bibtex_input", "english_dialect", "humanizer_mode"]}
        )
        db.session.add(new_task)
        db.session.commit()

        process_humanization_task.delay(
            task_id=new_task.id
        )
        
        return jsonify({"status": "processing_started", "transaction_id": new_task.external_reference})

    except AllModelsFailedError as e:
        db.session.rollback()
        flash(e.message, "danger")
        return jsonify({"status": "error", "redirect": url_for('main.index')}), 500
    except Exception as e:
        db.session.rollback()
        logger.error(f"An unexpected error occurred during humanization for user {user.id}: {e}", exc_info=True)
        flash("An unexpected error occurred. Please try again.", "danger")
        return jsonify({"status": "error", "redirect": url_for('main.index')}), 500

def _handle_anonymous_humanization():
    data = request.form
    input_text, phone_number = data.get("input_text", "").strip(), data.get("phone_number", "").strip()
    word_count = count_words(input_text)
    humanizer_mode = data.get("humanizer_mode", "standard")

    if not all([input_text, phone_number]):
        return jsonify({"status": "error", "message": "Text and phone number are required."}), 400

    if humanizer_mode != 'document_formatter' and word_count > MAX_WORDS:
        return jsonify({"status": "error", "message": f"Input exceeds the maximum of {MAX_WORDS} words."}), 413

    if not KENYAN_PHONE_REGEX.match(phone_number):
        return jsonify({"status": "error", "message": "Invalid Kenyan phone number format."}), 400
    price = ESSAY_PRICE_PER_UNIT + 10 if humanizer_mode == 'essay' else PRICE_PER_UNIT + 10
    amount_to_charge = round((word_count / WORDS_PER_UNIT) * price, 2)
    transaction_ref = f"HUMANIZE-{uuid.uuid4()}"
    new_task = Task(external_reference=transaction_ref, input_text=input_text, phone_number=phone_number,
                    amount_charged=amount_to_charge, status='payment_pending', user_id=None,
                    diff_data={key: data.get(key, "") for key in ["output_format", "bibtex_input", "english_dialect", "humanizer_mode"]})
    db.session.add(new_task)
    db.session.commit()
    try:
        initiate_payhero_payment(amount_to_charge, phone_number, transaction_ref)
        return jsonify({"status": "payment_initiated", "transaction_id": transaction_ref})
    except Exception as e:
        new_task.status, new_task.error_message = 'payment_failed', str(e)
        db.session.commit()
        logger.error(f"Anonymous payment failed: {e}", exc_info=True)
        return jsonify({"status": "error", "message": "Payment initiation failed."}), 500

@main_bp.route("/payhero-callback", methods=["POST"])
def payhero_callback():
    payload = request.get_json()
    logger.info(f"Received callback: {payload}")
    if not payload: return jsonify({'status': 'error', 'message': 'Invalid payload'}), 400
    response = payload.get('response', {})
    transaction_ref, status = response.get('ExternalReference'), response.get('Status')
    amount = float(response.get('Amount', 0))
    if not transaction_ref: return jsonify({'status': 'success'}), 200
    if transaction_ref.startswith("TOPUP-"):
        task = Task.query.filter_by(external_reference=transaction_ref, status='topup_pending').first()
        if task and task.user:
            if status == 'Success':
                user = task.user
                user.wallet.balance += amount
                db.session.add(Transaction(wallet_id=user.wallet.id, amount=amount,
                                           transaction_type='top-up', description=f"M-Pesa top-up from {task.phone_number}"))
                task.status = 'topup_complete'
            else:
                task.status = 'topup_failed'
                task.error_message = response.get('ResultDesc', 'Top-up was cancelled or failed.')
            db.session.commit()
            logger.info(f"Top-up task {task.id} status updated to {task.status}.")
    elif transaction_ref.startswith("HUMANIZE-"):
        task = Task.query.filter_by(external_reference=transaction_ref, status='payment_pending').first()
        if task:
            if status == 'Success':
                task.status = 'payment_successful'
                db.session.commit()
                app = current_app._get_current_object()
                from .tasks import process_humanization_in_background
                threading.Thread(target=process_humanization_in_background, args=(app, task.id)).start()
                logger.info(f"Anonymous task {task.id} payment successful.")
            else:
                task.status = 'payment_failed'
                task.error_message = response.get('ResultDesc', 'Payment was cancelled or failed.')
                db.session.commit()
                logger.info(f"Anonymous task {task.id} payment failed.")
    return jsonify({'status': 'success'}), 200

@main_bp.route("/get_humanization_status/<transaction_id>", methods=["GET"])
def get_humanization_status(transaction_id):
    task = Task.query.filter_by(external_reference=transaction_id).first_or_404()
    if task.user_id and ('user' not in session or task.user_id != session['user']['uid']):
        abort(403)
    
    response_data = {"status": task.status, "message": task.error_message or "Processing..."}

    if task.status == 'humanization_complete':
        diff_data = task.diff_data or {}
        output_format = diff_data.get("output_format", "plain_text")
        
        if output_format != 'plain_text':
            response_data.update({
                "status": "download_ready", 
                "download_url": url_for('main.download_formatted_doc', transaction_id=task.external_reference)
            })
        else:
            word_diff = diff_data.get('word_diff')
            html_output = ""
            if isinstance(word_diff, list):
                final_html_parts = []
                current_paragraph_spans = []
                for d in word_diff:
                    text_segment = d.get("text", "")
                    css_class = "word-added" if d.get("type") == "added" else "word-equal"
                    parts = text_segment.split("[[PARAGRAPH_BREAK]]")
                    if parts[0].strip():
                        current_paragraph_spans.append(f'<span class="{css_class}">{parts[0].strip()}</span>')
                    if len(parts) > 1:
                        for i in range(1, len(parts)):
                            if current_paragraph_spans:
                                final_html_parts.append(f'<p style="text-indent: 0;">{" ".join(current_paragraph_spans)}</p>')
                            current_paragraph_spans = []
                            if parts[i].strip():
                                current_paragraph_spans.append(f'<span class="{css_class}">{parts[i].strip()}</span>')
                if current_paragraph_spans:
                    final_html_parts.append(f'<p style="text-indent: 0;">{" ".join(current_paragraph_spans)}</p>')
                html_output = "".join(final_html_parts)
            else:
                html_output = task.humanized_output.replace('\n', '<br>')

            response_data.update({
                "status": "complete", 
                "humanized_output": html_output,
                "output_word_count": task.output_word_count, 
                "diff_data": diff_data
            })
            
    return jsonify(response_data)

@main_bp.route("/download_formatted_doc/<transaction_id>", methods=["GET"])
def download_formatted_doc(transaction_id):
    task = Task.query.filter_by(external_reference=transaction_id).first_or_404()
    if task.user_id and ('user' not in session or task.user_id != session['user']['uid']):
        abort(403)
    if task.status != 'humanization_complete':
        flash("Document is not yet ready for download.", "danger")
        return redirect(url_for('main.index'))

    output_dir = "/tmp" if os.environ.get("FLASK_ENV") == "production" else "output_docs"
    os.makedirs(output_dir, exist_ok=True)

    diff_data = task.diff_data or {}
    title = diff_data.get('generated_title') or "Processed document"
    filename = f"{sanitize_filename(title)}.docx"
    filepath = os.path.join(output_dir, filename)

    metadata = {'title': title, 'author': "Name", 'date': datetime.now(NAIROBI_TZ).strftime("%d %B %Y")}
    references = diff_data.get('formatted_references', [])

    if create_formatted_document(diff_data.get('output_format'), task.humanized_output, filepath, metadata, references):
        return send_file(filepath, as_attachment=True)

    flash("Failed to generate the document.", "danger")
    return redirect(url_for('main.index'))

from flask import render_template, abort, g, url_for
from markdown import markdown
# Ensure other necessary imports like Task, User, logger, etc., are present

@main_bp.route("/task/details/<transaction_id>")
@auth_required
def view_task_details(transaction_id):
    logger.info(f"--- Entering view_task_details for transaction_id: {transaction_id} ---")
    task = Task.query.filter_by(external_reference=transaction_id).first_or_404()
    user = User.query.get_or_404(session['user']['uid'])

    if not user.is_admin and (task.user_id != user.id):
        logger.warning(f"User {user.id} attempting to access unauthorized task {task.id}.")
        abort(403)

    # Fetch all necessary data for the template header and page
    wallet_balance = user.wallet.balance if user.wallet else 0
    username = user.name
    back_url = url_for('main.transaction_history')
    diff_data = task.diff_data or {}
    output_format = diff_data.get('output_format', 'plain_text')
    humanizer_mode = diff_data.get('humanizer_mode', 'N/A').replace('_', ' ').title()

    # Handle tasks that are not yet complete
    if task.status != 'humanization_complete':
        flash("This task is not yet complete.", "info")
        return redirect(request.referrer or url_for('main.transaction_history'))

    # Handle formatted documents by redirecting to download
    if output_format != 'plain_text':
        return redirect(url_for('main.download_formatted_doc', transaction_id=task.external_reference))

    # ✅ FIX: Prepare the HTML output with proper paragraph formatting
    html_output = "<p>Error: Could not render task output.</p>"
    try:
        if task.humanized_output:
            # This ensures that text separated by blank lines becomes distinct paragraphs
            paragraphs = task.humanized_output.strip().split('\n\n')
            html_output = "".join(f"<p>{p.strip()}</p>" for p in paragraphs if p.strip())
        else:
            html_output = "<p>Error: No output data found for this task.</p>"
    except Exception as e:
        logger.error(f"FATAL: Could not render HTML for task {transaction_id}. Error: {e}", exc_info=True)
        flash("An unexpected error occurred while trying to display this task.", "danger")
        return redirect(request.referrer or url_for('main.transaction_history'))

    # Render the final template with all the required data
    return render_template(
        "view_task.html", 
        task=task, 
        username=username, 
        html_output=html_output, 
        back_url=back_url, 
        humanizer_mode=humanizer_mode,
        wallet_balance=wallet_balance
    )

@main_bp.route("/instructions")
def instructions():
    user = User.query.get(session['user']['uid']) if 'user' in session else None
    username = get_display_name(user) if user else None
    wallet_balance = user.wallet.balance if user and user.wallet else 0
    is_admin = user.is_admin if user else False
    return render_template("instructions.html",
                           user=user,
                           username=username,
                           wallet_balance=wallet_balance,
                           is_admin=is_admin)

@main_bp.route("/admin/topup", methods=["POST"])
@auth_required
def admin_self_topup():
    user = User.query.get(session['user']['uid'])
    if not user or not user.is_admin:
        abort(403)

    data = request.get_json()
    amount = data.get("amount")

    if not isinstance(amount, (int, float)) or not (0 < amount <= 10000):
        return jsonify({"status": "error", "message": "Invalid amount. Must be between 0.01 and 10,000."}), 400

    user.wallet.balance += amount
    db.session.add(Transaction(
        wallet_id=user.wallet.id,
        amount=amount,
        transaction_type='admin_self_topup',
        description=f'Admin self top-up by {user.email}'
    ))
    db.session.commit()

    flash(f"Successfully added KSh {amount:.2f} to your wallet.", "success")
    return jsonify({"status": "success", "new_balance": user.wallet.balance})

@main_bp.route("/contact")
def contact():
    user = User.query.get(session['user']['uid']) if 'user' in session else None
    username = get_display_name(user) if user else None
    wallet_balance = user.wallet.balance if user and user.wallet else 0
    is_admin = user.is_admin if user else False

    return render_template("contact.html",
                           user=user,
                           whatsapp_number=WHATSAPP_NUMBER,
                           username=username,
                           wallet_balance=wallet_balance,
                           is_admin=is_admin)
  
@main_bp.route("/withdraw", methods=["POST"])
@auth_required
def withdraw_funds():
    from humanizer_app.models import ReferralWithdrawal, Transaction, User

    user = User.query.get_or_404(session['user']['uid'])

    # ✅ Calculate withdrawable amount (only referral commissions)
    withdrawable_amount = (
        db.session.query(func.sum(Transaction.amount))
        .filter(
            Transaction.wallet_id == user.wallet.id,
            Transaction.transaction_type == "referral_commission"
        )
        .scalar()
        or 0.0
    )

    # ✅ Validate amount
    if withdrawable_amount <= 0:
        flash("You have no referral commissions available for withdrawal.", "danger")
        return redirect(url_for("main.referrals_page"))

    if withdrawable_amount < 0.5:
        flash("Minimum withdrawal amount is KSh 0.50", "danger")
        return redirect(url_for("main.referrals_page"))

    # ✅ Deduct immediately from wallet balance
    if user.wallet.balance < withdrawable_amount:
        flash("Insufficient balance in your wallet to process withdrawal.", "danger")
        return redirect(url_for("main.referrals_page"))

    user.wallet.balance -= withdrawable_amount

    # ✅ Create Transaction record for history
    withdrawal_txn = Transaction(
        wallet_id=user.wallet.id,
        amount=-withdrawable_amount,
        transaction_type="withdrawal",
        status="pending",
        description="Referral withdrawal request submitted"
    )
    db.session.add(withdrawal_txn)

    # ✅ Create ReferralWithdrawal request for admin tracking
    withdrawal_request = ReferralWithdrawal(
        user_id=user.id,
        amount=withdrawable_amount,
        status="pending",
        transaction_id=None  # optional, can be linked later
    )
    db.session.add(withdrawal_request)

    db.session.commit()

    flash(f"Withdrawal request of KSh {withdrawable_amount:.2f} submitted successfully. Awaiting admin approval.", "success")
    return redirect(url_for("main.referrals_page"))


@main_bp.route("/get_conversations", methods=['GET'])
@auth_required
def get_conversations():
    user = User.query.get_or_404(session['user']['uid'])
    query = request.args.get('q', '')
    page = request.args.get('page', 1, type=int)
    per_page = 7

    base_query = Conversation.query.filter_by(user_id=user.id)

    if query:
        base_query = base_query.filter(Conversation.title.ilike(f'%{query}%'))

    paginated_convs = base_query.order_by(Conversation.created_at.desc()).paginate(
        page=page, per_page=per_page, error_out=False
    )

    conversations_data = [
        {"id": conv.id, "title": conv.title}
        for conv in paginated_convs.items
    ]

    return jsonify({
    'conversations': conversations_data,
    'total_pages': paginated_convs.pages,
    'current_page': paginated_convs.page,
    'has_next': paginated_convs.has_next,
    'has_prev': paginated_convs.has_prev
    })

@main_bp.route("/conversation/<int:conv_id>", methods=['GET'])
@auth_required
def get_conversation_messages(conv_id):
    user = User.query.get_or_404(session['user']['uid'])
    conversation = Conversation.query.filter_by(id=conv_id, user_id=user.id).first_or_404()
    messages = ChatMessage.query.filter_by(conversation_id=conversation.id).order_by(ChatMessage.created_at.asc()).all()
    message_list = [{"role": msg.role, "content": msg.content} for msg in messages]
    # +++ START OF MODIFIED CODE +++
    return jsonify({
        "messages": message_list,
        "model_key": conversation.model_key
    })
    # +++ END OF MODIFIED CODE +++

@main_bp.route("/conversation/<int:conv_id>", methods=['DELETE'])
@auth_required
def delete_conversation(conv_id):
    user = User.query.get_or_404(session['user']['uid'])
    conversation = Conversation.query.filter_by(id=conv_id, user_id=user.id).first()
    if conversation:
        db.session.delete(conversation)
        db.session.commit()
        return jsonify({"status": "success", "message": "Conversation deleted."}), 200
    return jsonify({"status": "error", "message": "Conversation not found or you do not have permission to delete it."}), 404

@main_bp.route('/extract_text_from_file', methods=['POST'])
@auth_required
def extract_text_from_file():
    if not check_rate_limit():
        return jsonify({"error": "You are uploading too frequently. Please wait a moment."}), 429

    if 'file' not in request.files:
        return jsonify({"error": "No file part"}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No selected file"}), 400

    if not is_extension_allowed(file.filename):
        return jsonify({"error": "This file type is not supported or is blocked for security reasons."}), 400

    filename = secure_filename(file.filename)
    file_bytes = file.read()
    file.seek(0)

    try:
        if filename.lower().endswith('.zip'):
            combined_text = f"--- Contents of {filename} ---\n\n"
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                for file_info in z.infolist():
                    if '.' in file_info.filename and file_info.filename.rsplit('.', 1)[1].lower() in BLOCKED_EXTENSIONS:
                        ext = file_info.filename.rsplit('.', 1)[1].lower()
                        return jsonify({"error": f"Upload failed. Executable file type (.{ext}) found inside ZIP."}), 400

                for file_info in z.infolist():
                    if file_info.is_dir():
                        continue
                    with z.open(file_info) as f:
                        file_content_bytes = f.read()
                        doc_text = f"\n--- File: {file_info.filename} ---\n"

                        if file_info.filename.lower().endswith('.docx'):
                            if b"EncryptedPackage" in file_content_bytes:
                                doc_text += "[File is encrypted and cannot be read]\n"
                            else:
                                doc = docx.Document(io.BytesIO(file_content_bytes))
                                doc_text += "".join(p.text + "\n" for p in doc.paragraphs)

                        elif file_info.filename.lower().endswith('.pdf'):
                            pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content_bytes))
                            if pdf_reader.is_encrypted:
                                doc_text += "[File is encrypted and cannot be read]\n"
                            else:
                                doc_text += "".join(page.extract_text() + "\n" for page in pdf_reader.pages)

                        else: 
                            doc_text += file_content_bytes.decode('utf-8', errors='ignore')

                        combined_text += doc_text
            return jsonify({"text": combined_text})

        elif filename.lower().endswith('.docx'):
            if b"EncryptedPackage" in file_bytes:
                return jsonify({"error": "This DOCX file is encrypted and cannot be processed."}), 400
            doc = docx.Document(io.BytesIO(file_bytes))
            extracted_text = "".join(p.text + "\n" for p in doc.paragraphs)
            return jsonify({"text": extracted_text})

        elif filename.lower().endswith('.pdf'):
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
            if pdf_reader.is_encrypted:
                return jsonify({"error": "This PDF file is encrypted and cannot be processed."}), 400
            extracted_text = "".join(page.extract_text() + "\n" for page in pdf_reader.pages)
            return jsonify({"text": extracted_text})

        elif filename.lower().endswith('.txt'):
            extracted_text = file_bytes.decode('utf-8', errors='ignore')
            return jsonify({"text": extracted_text})

        elif filename.rsplit('.', 1)[1].lower() in {'png', 'jpg', 'jpeg', 'webp', 'gif'}:
            return jsonify({"text": f"[Image '{filename}' attached]"})

    except Exception as e:
        logger.error(f"Error processing file {filename}: {e}", exc_info=True)
        return jsonify({"error": "Failed to process file. It may be corrupt or in an unsupported format."}), 500

    return jsonify({"error": "File could not be processed."}), 500

@main_bp.route('/conversations/clear_all', methods=['POST'])
@auth_required
def clear_all_conversations():
    try:
        user = User.query.get_or_404(session['user']['uid'])
        
        conversation_ids = [conv.id for conv in Conversation.query.filter_by(user_id=user.id).all()]

        if conversation_ids:
            ChatMessage.query.filter(ChatMessage.conversation_id.in_(conversation_ids)).delete(synchronize_session=False)
            
            Conversation.query.filter_by(user_id=user.id).delete(synchronize_session=False)
        
        db.session.commit()
        
        return jsonify({'message': 'All conversations cleared successfully.'}), 200
    except Exception as e:
        db.session.rollback()
        logger.error(f"Error clearing all conversations for user {user.id}: {e}", exc_info=True)
        return jsonify({'error': 'An internal error occurred while clearing history.'}), 500
