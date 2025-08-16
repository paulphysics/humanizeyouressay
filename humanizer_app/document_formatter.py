"""
This module provides functionality to create formatted .docx documents
in various academic styles (APA, MLA, Harvard).
"""
import os
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# --- Constants ---
DEFAULT_FONT = 'Times New Roman'
DEFAULT_FONT_SIZE = Pt(12)
DOUBLE_SPACING = 2.0
INCH_MARGIN = 1.0
HALF_INCH_INDENT = 0.5

# --- DocumentStyler Class ---
class DocumentStyler:
    """A class to manage document styling for different formats."""

    def __init__(self, document, style_name):
        self.doc = document
        self.style_name = style_name
        self._apply_base_styles()

    def _apply_base_styles(self):
        """Applies universal styles like margins and default font."""
        for section in self.doc.sections:
            section.top_margin = Inches(INCH_MARGIN)
            section.bottom_margin = Inches(INCH_MARGIN)
            section.left_margin = Inches(INCH_MARGIN)
            section.right_margin = Inches(INCH_MARGIN)

        style = self.doc.styles['Normal']
        style.font.name = DEFAULT_FONT
        style.font.size = DEFAULT_FONT_SIZE
        p_format = style.paragraph_format
        p_format.line_spacing = DOUBLE_SPACING
        p_format.space_before = Pt(0)
        p_format.space_after = Pt(0)

    def add_header_with_page_number(self, author_last_name=""):
        """Adds a header with a page number."""
        header = self.doc.sections[0].header
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if author_last_name:
            p.add_run(f"{author_last_name} ")

        page_num_run = p.add_run()
        fld_char_begin = OxmlElement('w:fldChar')
        fld_char_begin.set(qn('w:fldCharType'), 'begin')
        instr_text = OxmlElement('w:instrText')
        instr_text.text = 'PAGE'
        fld_char_end = OxmlElement('w:fldChar')
        fld_char_end.set(qn('w:fldCharType'), 'end')

        page_num_run._r.append(fld_char_begin)
        page_num_run._r.append(instr_text)
        page_num_run._r.append(fld_char_end)


    def add_title_page(self, metadata):
        """Adds a title page or header based on the citation style."""
        if self.style_name == 'apa_7':
            self._add_apa_title_page(metadata)
        elif self.style_name == 'mla_9':
            # MLA does not have a separate title page.
            pass
        elif self.style_name == 'harvard':
            self._add_harvard_title_page(metadata)

    def add_body_content(self, content, metadata):
        """
        Adds the main body content, identifying and formatting subheadings
        based on Markdown syntax and word count.
        """
        if self.style_name in ['apa_7', 'harvard']:
            self.doc.add_page_break()

        # Add the appropriate title or header for the first page of content
        if self.style_name == 'mla_9':
            self._add_mla_header(metadata)
        else: # For APA and Harvard, repeat the title
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(metadata.get('title', 'Paper Title').replace('"', ''))
            if self.style_name == 'apa_7':
                run.bold = True
            p.paragraph_format.space_after = Pt(24)


        for paragraph_text in content.split('\n'):
            stripped_text = paragraph_text.strip()
            if not stripped_text:
                continue

            # --- THIS IS THE FIX: Improved Heading Detection ---
            is_markdown_heading = stripped_text.startswith(('#', '##', '###', '####'))
            ends_with_punctuation = bool(re.search(r'[.!?]$', stripped_text))
            is_title_case_heading = stripped_text.istitle() and not ends_with_punctuation
            is_short_line_heading = (len(stripped_text.split()) < 7 and not ends_with_punctuation)

            if is_markdown_heading:
                if stripped_text.startswith('#### '):
                    self._add_level_4_heading(stripped_text.lstrip('# ').strip())
                elif stripped_text.startswith('### '):
                    self._add_level_3_heading(stripped_text.lstrip('# ').strip())
                elif stripped_text.startswith('## '):
                    self._add_level_2_heading(stripped_text.lstrip('# ').strip())
                elif stripped_text.startswith('# '):
                    self._add_level_1_heading(stripped_text.lstrip('# ').strip())
            elif is_title_case_heading or is_short_line_heading:
                self._add_level_1_heading(stripped_text)
            else:
                p = self.doc.add_paragraph(stripped_text)
                p.paragraph_format.first_line_indent = Inches(HALF_INCH_INDENT)
            # --- END OF FIX ---

    def add_references_section(self, references_data):
        """Adds the references section."""
        if not references_data:
            return
            
        if self.style_name == 'mla_9':
            title = "Works Cited"
        elif self.style_name == 'harvard':
            title = "Reference List"
        else: # APA 7
            title = "References"
            
        self.doc.add_page_break()
        heading = self.doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(title)
        if self.style_name == 'apa_7':
            run.bold = True

        for entry in references_data:
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(HALF_INCH_INDENT)
            p.paragraph_format.first_line_indent = Inches(-HALF_INCH_INDENT)
            
            parts = re.split(r'(\*.*?\*)', entry)
            for part in parts:
                if part.startswith('*') and part.endswith('*'):
                    p.add_run(part[1:-1]).italic = True
                else:
                    p.add_run(part)

    def _add_apa_title_page(self, metadata):
        """Adds an APA 7 style title page."""
        self.doc.add_paragraph('\n' * 4)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(metadata.get('title', 'Paper Title').replace('"', '')).bold = True
        p.add_run(f"\n\n{metadata.get('author', 'Name')}")
        p.add_run(f"\n{metadata.get('institution', 'Institution')}")
        p.add_run(f"\n{metadata.get('course', 'Course')}")
        p.add_run(f"\n{metadata.get('instructor', 'Instructor')}")
        p.add_run(f"\n{metadata.get('date', 'Date')}")

    def _add_mla_header(self, metadata):
        """Adds the four-line header for MLA style."""
        p = self.doc.add_paragraph()
        p.add_run(f"{metadata.get('author', 'Name')}\n")
        p.add_run(f"{metadata.get('instructor', 'Instructor')}\n")
        p.add_run(f"{metadata.get('course', 'Course')}\n")
        p.add_run(f"{metadata.get('date', 'Date')}\n\n")

        title_p = self.doc.add_paragraph(metadata.get('title', 'Paper Title').replace('"', ''))
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _add_harvard_title_page(self, metadata):
        """Adds a Harvard style title page."""
        p = self.doc.add_paragraph('\n' * 4)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
        title_run = p.add_run(metadata.get('title', 'Paper Title').replace('"', '').upper())
        title_run.bold = True
        title_run.font.size = Pt(12)
    
        spacing_and_author = '\n' * 5 + f"by\n\n{metadata.get('author', 'Name')}"
        spacing_and_author += '\n' * 2 + f"{metadata.get('course', 'Course')}\n"
        spacing_and_author += f"{metadata.get('institution', 'Institution')}\n"
        spacing_and_author += f"Date: {metadata.get('date', 'Date')}"
        
        p.add_run(spacing_and_author)

    def _add_level_1_heading(self, text):
        """Adds a Level 1 heading formatted for the current style."""
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        if self.style_name == 'apa_7':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run.bold = True
        elif self.style_name == 'mla_9':
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run.bold = True
        elif self.style_name == 'harvard':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run.bold = True

    def _add_level_2_heading(self, text):
        """Adds a Level 2 heading formatted for the current style."""
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        if self.style_name == 'apa_7':
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run.bold = True
        elif self.style_name == 'mla_9':
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run.italic = True
        elif self.style_name == 'harvard':
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run.bold = True

    def _add_level_3_heading(self, text):
        """Adds a Level 3 heading formatted for the current style."""
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        if self.style_name == 'apa_7':
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run.bold = True
            run.italic = True
        elif self.style_name == 'mla_9':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run.bold = True
        elif self.style_name == 'harvard':
            p.paragraph_format.left_indent = Inches(HALF_INCH_INDENT)
            run.bold = True
            run.italic = True

    def _add_level_4_heading(self, text):
        """Adds a Level 4 heading formatted for the current style."""
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        if self.style_name == 'apa_7':
            p.paragraph_format.left_indent = Inches(HALF_INCH_INDENT)
            run.bold = True
            p.add_run(".")
        elif self.style_name == 'mla_9':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run.italic = True
        elif self.style_name == 'harvard':
            p.paragraph_format.left_indent = Inches(HALF_INCH_INDENT)
            run.italic = True


def create_formatted_document(format_type: str, humanized_content: str, output_filename: str, metadata: dict, references_data: list = None):
    document = Document()
    
    core_props = document.core_properties
    core_props.author = metadata.get('author', 'User')
    core_props.title = metadata.get('title', 'My Document')
    core_props.comments = "Written by me!"
    
    styler = DocumentStyler(document, format_type)
    
    if format_type in ['apa_7', 'harvard']:
        styler.add_header_with_page_number()
    elif format_type == 'mla_9':
        styler.add_header_with_page_number(author_last_name=metadata.get('author', 'Name'))

    styler.add_title_page(metadata)
    styler.add_body_content(humanized_content, metadata)
    styler.add_references_section(references_data)

    try:
        document.save(output_filename)
        print(f"Successfully created {output_filename}")
        return True
    except Exception as e:
        print(f"Error saving document: {e}")
        return False
